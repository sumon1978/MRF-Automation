# FINAL VERSION: CC_MOD + MRF Edit upload/download workflow
import streamlit as st
from openpyxl import load_workbook
import pandas as pd
import os
import re
import tempfile
import io
import zipfile


# ============================================================
# STREAMLIT PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="MRF Automation",
    page_icon="📊",
    layout="wide"
)


# ============================================================
# PROJECT DASHBOARD MENU
# ============================================================

def project_dashboard():

    st.markdown("""
    <style>
    .title-box {
        background: linear-gradient(135deg,#0f172a,#2563eb);
        padding:30px;
        border-radius:20px;
        text-align:center;
        color:white;
        margin-bottom:25px;
    }
    .project-card {
        background:white;
        padding:18px;
        border-radius:15px;
        border:1px solid #e2e8f0;
        box-shadow:0 8px 20px rgba(0,0,0,.08);
        text-align:center;
        font-size:20px;
        font-weight:700;
        margin-bottom:12px;
    }
    </style>
    """, unsafe_allow_html=True)

    st.markdown("""
    <div class="title-box">
        <h1>MRF Automation</h1>
    </div>
    """, unsafe_allow_html=True)

    st.subheader("Project wise")

    projects = ["CC_MOD", "New BTS", "5G", "TDD", "Others"]

    for project in projects:

        st.markdown(
            f'<div class="project-card">{project}</div>',
            unsafe_allow_html=True
        )

        if st.button(
            f"Open {project}",
            key=f"project_{project}",
            use_container_width=True
        ):
            st.session_state["selected_project"] = project

    return st.session_state.get("selected_project")


if "selected_project" not in st.session_state:
    st.session_state["selected_project"] = None

selected_project = project_dashboard()

if selected_project in ["New BTS", "5G", "TDD", "Others"]:
    st.markdown("""
    <div style="
        background:#fff7ed;
        border:1px solid #fed7aa;
        padding:35px;
        border-radius:18px;
        text-align:center;">
        <h2>🚧 Under Construction</h2>
        <p>This project module is currently under development.</p>
    </div>
    """, unsafe_allow_html=True)
    st.stop()

# CC_MOD selected or first load continues to original MRF logic


# ============================================================
# PROFESSIONAL WEB UI
# ============================================================

st.markdown("""
<style>
    .stApp {
        background:
            radial-gradient(circle at 10% 0%, rgba(37,99,235,.10), transparent 28%),
            radial-gradient(circle at 90% 10%, rgba(14,165,233,.08), transparent 25%),
            #f6f8fc;
    }
    .block-container {
        max-width: 1280px;
        padding-top: 1.6rem;
        padding-bottom: 3rem;
    }
    header[data-testid="stHeader"] {
        background: rgba(255,255,255,.72);
        backdrop-filter: blur(10px);
    }
    .hero {
        padding: 34px 38px;
        border-radius: 24px;
        background: linear-gradient(135deg, #0f172a 0%, #172554 48%, #1d4ed8 100%);
        box-shadow: 0 18px 45px rgba(15,23,42,.16);
        margin-bottom: 22px;
        color: white;
    }
    .hero-badge {
        display: inline-block;
        padding: 7px 12px;
        border: 1px solid rgba(255,255,255,.20);
        border-radius: 999px;
        background: rgba(255,255,255,.10);
        font-size: 12px;
        font-weight: 700;
        letter-spacing: .06em;
        text-transform: uppercase;
        margin-bottom: 14px;
    }
    .hero h1 {
        margin: 0 0 10px 0;
        color: white;
        font-size: 38px;
        line-height: 1.12;
    }
    .hero p {
        max-width: 760px;
        margin: 0;
        color: #dbeafe;
        font-size: 16px;
        line-height: 1.7;
    }
    .section-title {
        font-size: 21px;
        font-weight: 800;
        color: #0f172a;
        margin: 8px 0 4px;
    }
    .section-copy {
        color: #64748b;
        margin: 0 0 14px;
        font-size: 14px;
    }
    div[data-testid="stFileUploader"] {
        background: #ffffff;
        border: 1px solid #e2e8f0;
        border-radius: 18px;
        padding: 12px;
        min-height: 185px;
        box-shadow: 0 8px 24px rgba(15,23,42,.05);
    }
    div[data-testid="stFileUploader"] section {
        border-radius: 13px;
        border: 1.5px dashed #bfdbfe;
        background: #f8fbff;
    }
    div.stButton > button {
        min-height: 52px;
        border-radius: 14px;
        font-weight: 800;
        border: 0;
        box-shadow: 0 10px 24px rgba(37,99,235,.18);
    }
    div[data-testid="stDownloadButton"] > button {
        min-height: 48px;
        border-radius: 13px;
        font-weight: 700;
        border: 1px solid #dbe3ef;
        background: white;
        color: #0f172a;
    }
    div[data-testid="stAlert"] {
        border-radius: 14px;
    }
    hr {
        border-color: #e5e7eb !important;
        margin: 1.5rem 0 !important;
    }
    .steps {
        display: grid;
        grid-template-columns: repeat(3, 1fr);
        gap: 12px;
        margin: 8px 0 22px;
    }
    .step {
        background: white;
        border: 1px solid #e2e8f0;
        border-radius: 15px;
        padding: 14px 16px;
        color: #475569;
        box-shadow: 0 5px 18px rgba(15,23,42,.035);
    }
    .step b { color: #0f172a; }
    .footer-note {
        text-align: center;
        color: #94a3b8;
        font-size: 12px;
        margin-top: 30px;
    }
    @media (max-width: 800px) {
        .block-container { padding-left: 1rem; padding-right: 1rem; }
        .hero { padding: 26px 22px; border-radius: 19px; }
        .hero h1 { font-size: 30px; }
        .steps { grid-template-columns: 1fr; }
    }
</style>
""", unsafe_allow_html=True)

st.markdown("""
<div class="hero">
    <div class="hero-badge">MRF Processing Workspace</div>
    <h1>MRF Automation</h1>
    <p>
        Generate site-wise MRF Excel files, apply item-code rules automatically,
        and download the completed output individually or as one ZIP package.
    </p>
</div>
""", unsafe_allow_html=True)

st.markdown("""
<div class="steps">
    <div class="step"><b>01 · Upload</b><br>Provide the three required Excel files.</div>
    <div class="step"><b>02 · Process</b><br>Generate MRFs and apply item-code rules.</div>
    <div class="step"><b>03 · Download</b><br>Download individual files or one ZIP.</div>
</div>
""", unsafe_allow_html=True)

st.markdown('<div class="section-title">Upload source files</div>', unsafe_allow_html=True)
st.markdown(
    '<div class="section-copy">All three Excel files are required before processing can begin.</div>',
    unsafe_allow_html=True
)



# ============================================================
# MRF EDIT FINAL WORKFLOW FUNCTIONS
# ============================================================

def norm(v):
    if pd.isna(v):
        return ""

    return str(v).strip().upper().replace(" ", "")



def load_edit_rules(edit_rule_file):

    df_rules = pd.read_excel(
        edit_rule_file,
        sheet_name="RULES",
        header=None,
        dtype=object
    ).fillna("")

    rules = []

    # Rule file:
    # B = Old Material
    # C = New Material
    # D = Plant
    # E = Storage Location
    # H = New Description

    for _, row in df_rules.iterrows():

        old_material = norm(row.iloc[1] if len(row) > 1 else "")

        if not old_material:
            continue

        rules.append({
            "old_material": old_material,
            "new_material": str(row.iloc[2]).strip() if len(row) > 2 else "",
            "plant": norm(row.iloc[3] if len(row) > 3 else ""),
            "storage": norm(row.iloc[4] if len(row) > 4 else ""),
            "description": str(row.iloc[7]).strip() if len(row) > 7 else ""
        })

    return rules


def process_mrf_edit(edit_rule_file, mrf_files):

    rules = load_edit_rules(edit_rule_file)

    edited_files = []

    for uploaded in mrf_files:

        temp_dir = tempfile.mkdtemp()
        input_path = os.path.join(temp_dir, uploaded.name)

        with open(input_path, "wb") as f:
            f.write(uploaded.getbuffer())

        wb = load_workbook(input_path)

        updated = 0

        for ws in wb.worksheets:

            for r in range(1, ws.max_row + 1):

                # MRF columns:
                # C = Material Code
                # D = Plant
                # E = Storage Location
                # G = Material Description

                material = norm(ws[f"C{r}"].value)
                plant = norm(ws[f"D{r}"].value)
                storage = norm(ws[f"E{r}"].value)

                if not material:
                    continue

                for rule in rules:

                    if material != rule["old_material"]:
                        continue

                    if rule["plant"] and plant != rule["plant"]:
                        continue

                    if rule["storage"] and storage != rule["storage"]:
                        continue

                    ws[f"C{r}"] = rule["new_material"]

                    if rule["description"]:
                        ws[f"G{r}"] = rule["description"]

                    updated += 1
                    break

        output_path = input_path.replace(".xlsx", "_EDITED.xlsx")

        wb.save(output_path)
        wb.close()

        edited_files.append(output_path)

    return edited_files



# ============================================================
# CC_MOD MODULE MENU
# ============================================================

if "ccmod_module" not in st.session_state:
    st.session_state["ccmod_module"] = "MRF Make"

st.markdown("### CC_MOD")

m1, m2 = st.columns(2)

with m1:
    if st.button("📄 MRF Make", use_container_width=True):
        st.session_state["ccmod_module"] = "MRF Make"

with m2:
    if st.button("✏️ MRF Edit", use_container_width=True):
        st.session_state["ccmod_module"] = "MRF Edit"


if st.session_state["ccmod_module"] == "MRF Edit":

    st.markdown("## MRF Edit")

    edit_rule_file = st.file_uploader(
        "Upload Item code _edit _site wise_final.xlsx",
        type=["xlsx"],
        key="edit_rule_file"
    )

    mrf_files = st.file_uploader(
        "Upload MRF Files (Multiple .xlsx)",
        type=["xlsx"],
        accept_multiple_files=True,
        key="edit_mrf_files"
    )

    st.info("""
    Workflow:

    1. Upload Item code _edit _site wise_final.xlsx
    2. Upload MRF Files
    3. Apply Edit Rules
       - Site matching
       - Old Material → New Material
       - Description update
    4. Download Edited MRF Files
    5. Download All Edited MRF Files (ZIP)
    """)

    edited_upload = st.file_uploader(
        "Upload Edited MRF Files (Optional Re-upload)",
        type=["xlsx"],
        accept_multiple_files=True,
        key="edited_mrf_files"
    )

    if st.button("🚀 Apply Edit Rules", use_container_width=True):

        if edit_rule_file and mrf_files:

            with st.spinner("Applying edit rules..."):

                edited_files = process_mrf_edit(
                    edit_rule_file,
                    mrf_files
                )

            st.success(
                f"Edited files completed: {len(edited_files)}"
            )

            st.markdown("### Download Edited MRF Files")

            for file_path in edited_files:

                with open(file_path, "rb") as f:
                    st.download_button(
                        "Download " + os.path.basename(file_path),
                        f.read(),
                        file_name=os.path.basename(file_path),
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )

            zip_buffer = io.BytesIO()

            with zipfile.ZipFile(
                zip_buffer,
                "w",
                zipfile.ZIP_DEFLATED
            ) as z:

                for file_path in edited_files:
                    z.write(
                        file_path,
                        os.path.basename(file_path)
                    )

            zip_buffer.seek(0)

            st.download_button(
                "📦 Download All Edited MRF Files (ZIP)",
                zip_buffer,
                file_name="Edited_MRF_All_Files.zip",
                mime="application/zip",
                use_container_width=True
            )

        else:
            st.error("Please upload rule file and MRF files.")

    st.stop()


# ============================================================
# FILE UPLOAD
# ============================================================

col1, col2, col3 = st.columns(3, gap="medium")

with col1:
    template_uploaded = st.file_uploader(
        "Upload MRF Template",
        type=["xlsx"],
        key="template"
    )

with col2:
    from_uploaded = st.file_uploader(
        "Upload From Solution Data",
        type=["xlsx"],
        key="from_solution"
    )

with col3:
    rules_uploaded = st.file_uploader(
        "Upload Item Code Rules File",
        type=["xlsx"],
        key="rules"
    )

st.caption("Accepted format: .xlsx • Files are processed temporarily during this session.")
st.divider()


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def clean_list(data):
    return [
        str(v).strip()
        for v in data
        if str(v).strip() != ""
        and str(v).lower() != "nan"
    ]


def safe_num(val):
    try:
        return int(float(val))
    except Exception:
        return 0


def safe_col(data, index):
    if len(data.columns) > index:
        return data.iloc[:, index].tolist()
    return []


def check_model(value, models):
    text = str(value).upper().replace(" ", "")
    return any(model in text for model in models)


def check_band(value, band):
    return band in str(value).upper()



# ============================================================
# LOAD ITEM CODE RULES
# ============================================================

def load_item_code_rules(rules_path):
    df_rules = pd.read_excel(
        rules_path,
        sheet_name="RULES",
        dtype=object
    ).fillna("")

    df_rules = df_rules.infer_objects(copy=False)

    exact_map = {}
    smart_map = {}

    for i in range(len(df_rules)):
        key = norm(df_rules.iloc[i, 0])

        if not key:
            continue

        values = df_rules.iloc[i, 1:6].tolist()
        exact_map[key] = values

        if "B1" in key or "B3" in key or "B8" in key:
            smart_map[key] = values

    return exact_map, smart_map


def smart_match(g, smart_map):
    for key, values in smart_map.items():
        rule = key.replace(" ", "")

        if "B1" in rule:
            b_part = "B1"
            item_part = rule.replace("B1", "")
        elif "B3" in rule:
            b_part = "B3"
            item_part = rule.replace("B3", "")
        elif "B8" in rule:
            b_part = "B8"
            item_part = rule.replace("B8", "")
        else:
            continue

        if b_part not in g:
            continue

        codes = item_part.split("/")

        for code in codes:
            if code and code in g:
                return values

    return None


def write_rule_values(ws, row, values, fallback_item=None):
    ws[f"C{row}"] = values[0] if len(values) > 0 else ""
    ws[f"D{row}"] = values[1] if len(values) > 1 else ""
    ws[f"E{row}"] = values[2] if len(values) > 2 else ""
    ws[f"F{row}"] = values[3] if len(values) > 3 else ""

    if len(values) > 4:
        ws[f"G{row}"] = values[4]
    elif fallback_item is not None:
        ws[f"G{row}"] = fallback_item


def apply_item_codes(wb, exact_map, smart_map):
    updated = 0

    for ws in wb.worksheets:

        for r in range(1, ws.max_row + 1):
            g = norm(ws[f"G{r}"].value)

            if not g:
                continue

            values = exact_map.get(g)

            if values is None:
                values = smart_match(g, smart_map)

            if values is not None:
                write_rule_values(
                    ws,
                    r,
                    values,
                    ws[f"G{r}"].value
                )
                updated += 1

        # Extra rules created by 4480 / 4499 logic
        for row, item_name in [
            (31, "Circular Power Connector"),
            (32, "Dual ERS heavy bracket")
        ]:
            if ws[f"I{row}"].value == 1:
                key = norm(item_name)

                if key in exact_map:
                    write_rule_values(
                        ws,
                        row,
                        exact_map[key],
                        item_name
                    )
                    updated += 1
                else:
                    ws[f"G{row}"] = item_name

    return updated



# ============================================================
# MRF EDIT MODULE FUNCTIONS
# ============================================================





# ============================================================
# MAIN MRF PROCESSING FUNCTION
# ============================================================

def create_mrf_files(template_file, from_file, rules_file):

    created_files = []
    temp_dir = tempfile.mkdtemp()

    template_path = os.path.join(temp_dir, "MRF_Sample.xlsx")
    from_path = os.path.join(temp_dir, "from_solution_file.xlsx")
    rules_path = os.path.join(temp_dir, "item_code_rules.xlsx")

    output_folder = os.path.join(temp_dir, "MRF_Make_output_file")
    os.makedirs(output_folder, exist_ok=True)

    # Save uploaded files
    with open(template_path, "wb") as f:
        f.write(template_file.getbuffer())

    with open(from_path, "wb") as f:
        f.write(from_file.getbuffer())

    with open(rules_path, "wb") as f:
        f.write(rules_file.getbuffer())

    # Read source and rules
    df = pd.read_excel(from_path)
    df.iloc[:, 0] = (
        df.iloc[:, 0]
        .fillna("")
        .astype(str)
        .str.strip()
    )

    site_ids = df.iloc[:, 0].unique()
    exact_map, smart_map = load_item_code_rules(rules_path)

    # ========================================================
    # PROCESS EACH SITE
    # ========================================================

    for site in site_ids:

        if not site:
            continue

        matched_rows = df[
            df.iloc[:, 0]
            .astype(str)
            .str.strip()
            == str(site).strip()
        ]

        if matched_rows.empty:
            continue

        wb = load_workbook(template_path)
        ws = wb.active

        # ----------------------------------------------------
        # BASIC VALUES
        # ----------------------------------------------------

        val_P = matched_rows.iloc[0, 15] if len(matched_rows.columns) > 15 else ""
        val_Q = matched_rows.iloc[0, 16] if len(matched_rows.columns) > 16 else ""
        val_R = matched_rows.iloc[0, 17] if len(matched_rows.columns) > 17 else ""
        val_S = matched_rows.iloc[0, 18] if len(matched_rows.columns) > 18 else ""
        val_T = matched_rows.iloc[0, 19] if len(matched_rows.columns) > 19 else ""

        ws["G8"] = val_P
        ws["G9"] = val_Q
        ws["G10"] = site
        ws["G11"] = val_R

        ws["J8"] = val_T
        ws["J9"] = val_S
        ws["J10"] = val_T
        ws["J11"] = val_S

        # ----------------------------------------------------
        # COUNTERS
        # ----------------------------------------------------

        count_2219_b8 = 0
        count_4415_4428_b3 = 0
        count_2219_2217_b1 = 0
        count_4480_4499_b1_b3 = 0

        for _, row in matched_rows.iterrows():

            b8 = str(row.iloc[1]).upper()
            b1 = str(row.iloc[2]).upper()
            b3 = str(row.iloc[3]).upper()

            full_row = f"{b8} {b1} {b3}"

            # 2219 B8
            if check_model(full_row, ["2219"]) and check_band(b8, "B8"):
                count_2219_b8 += 1

            # 4415 / 4428 B3
            if check_model(full_row, ["4415", "4428"]) and check_band(b3, "B3"):
                count_4415_4428_b3 += 1

            # 2219 / 2217 B1
            model_text = full_row.replace(" ", "")

            if (
                ("2219/2217" in model_text or "2217/2219" in model_text)
                and "B1" in b1
            ):
                count_2219_2217_b1 += 1

            # 4480 / 4499 B1 B3
            if (
                check_model(full_row, ["4480", "4499"])
                and "B1" in full_row
                and "B3" in full_row
            ):
                count_4480_4499_b1_b3 += 1

        # ----------------------------------------------------
        # WRITE COUNTS
        # ----------------------------------------------------

        ws["G15"] = "2219 B8" if count_2219_b8 else ""
        ws["I15"] = count_2219_b8

        ws["G16"] = "4415/4428 B3" if count_4415_4428_b3 else ""
        ws["I16"] = count_4415_4428_b3

        ws["G17"] = "2219/2217 B1" if count_2219_2217_b1 else ""
        ws["I17"] = count_2219_2217_b1

        ws["G18"] = "4480/4499 B1 B3" if count_4480_4499_b1_b3 else ""
        ws["I18"] = count_4480_4499_b1_b3

        # ----------------------------------------------------
        # 4480 / 4499 EXTRA ITEMS
        # ----------------------------------------------------

        if count_4480_4499_b1_b3:
            ws["G31"] = "Circular Power Connector"
            ws["I31"] = 1

            ws["G32"] = "Dual ERS heavy bracket"
            ws["I32"] = 1
        else:
            ws["G31"] = ""
            ws["I31"] = ""
            ws["G32"] = ""
            ws["I32"] = ""

        # ----------------------------------------------------
        # GSM
        # ----------------------------------------------------

        gsm_total = sum(
            safe_num(v)
            for v in clean_list(safe_col(matched_rows, 4))
        )

        if gsm_total != 0:
            ws["G24"] = "GSM"
            ws["I24"] = gsm_total

            ws["G25"] = "Down Tilt H"
            ws["I25"] = gsm_total

            ws["G26"] = "SR:RET Control Cable (3GPP / AISG) 5 m"
            ws["I26"] = gsm_total
        else:
            for cell in ["G24", "I24", "G25", "I25", "G26", "I26"]:
                ws[cell] = ""

        # ----------------------------------------------------
        # SUM COLUMNS
        # ----------------------------------------------------

        item_rules = [
            (8,  "G19", "I19", "SFP"),
            (9,  "G20", "I20", "RRU Connector"),
            (10, "G21", "I21", "4.3 to 4.3 jumper"),
            (11, "G22", "I22", "4.3 to 7/16 jumper"),
            (12, "G23", "I23", "RRU pw cable"),
        ]

        for col_index, g_cell, i_cell, item_name in item_rules:
            total = sum(
                safe_num(v)
                for v in clean_list(safe_col(matched_rows, col_index))
            )
            ws[g_cell] = item_name
            ws[i_cell] = total if total else 0

        # ----------------------------------------------------
        # TOTAL QUANTITY
        # ----------------------------------------------------

        extra_minus = 1 if count_4480_4499_b1_b3 else 0

        total_quantity = (
            safe_num(ws["I15"].value)
            + safe_num(ws["I16"].value)
            + safe_num(ws["I17"].value)
            + safe_num(ws["I18"].value)
            - extra_minus
        )

        ws["I27"] = total_quantity if total_quantity != 0 else ""
        ws["I28"] = total_quantity if total_quantity != 0 else ""

        # ----------------------------------------------------
        # HIDE / SHOW ROWS
        # ----------------------------------------------------

        for r in range(15, 33):
            value = ws[f"I{r}"].value
            ws.row_dimensions[r].hidden = value in [0, "0", None, ""]

        # ----------------------------------------------------
        # APPLY ITEM CODES (SECOND SCRIPT)
        # ----------------------------------------------------

        apply_item_codes(wb, exact_map, smart_map)

        # ----------------------------------------------------
        # SUFFIX
        # ----------------------------------------------------

        suffix = "4K"

        if val_R:
            r_text = str(val_R).strip().upper()

            if "GPI" in r_text:
                suffix = "GPI"
            elif "CFVL" in r_text:
                suffix = "CFVL"
            else:
                suffix = r_text

        safe_site = re.sub(
            r'[\\/*?:"<>|]',
            "_",
            str(site)
        )

        output_path = os.path.join(
            output_folder,
            f"MRF_{safe_site}_{suffix}.xlsx"
        )

        wb.save(output_path)
        wb.close()

        created_files.append(output_path)

    return created_files


# ============================================================
# GENERATE BUTTON
# ============================================================

st.markdown('<div class="section-title">Generate output</div>', unsafe_allow_html=True)
st.markdown(
    '<div class="section-copy">Once all files are uploaded, start the automated MRF processing workflow.</div>',
    unsafe_allow_html=True
)

if st.button(
    "🚀 Generate MRF Files + Add Item Codes",
    type="primary",
    use_container_width=True
):

    if template_uploaded is None:
        st.error("❌ Please upload the MRF Template Excel file.")

    elif from_uploaded is None:
        st.error("❌ Please upload the From Solution Data Excel file.")

    elif rules_uploaded is None:
        st.error("❌ Please upload the Item Code Rules Excel file.")

    else:
        try:
            with st.spinner(
                "Generating MRF files and adding item codes... Please wait."
            ):
                created_files = create_mrf_files(
                    template_uploaded,
                    from_uploaded,
                    rules_uploaded
                )

            if not created_files:
                st.warning("⚠️ No MRF files were created.")

            else:
                st.success(
                    f"✅ Completed successfully! Total files created: "
                    f"{len(created_files)}"
                )

                st.divider()
                st.markdown('<div class="section-title">Generated MRF files</div>', unsafe_allow_html=True)
                st.caption("Your processed files are ready. Download them below.")

                for file_path in created_files:
                    file_name = os.path.basename(file_path)

                    with open(file_path, "rb") as f:
                        file_data = f.read()

                    st.download_button(
                        label=f"⬇️ Download {file_name}",
                        data=file_data,
                        file_name=file_name,
                        mime=(
                            "application/vnd.openxmlformats-officedocument."
                            "spreadsheetml.sheet"
                        ),
                        use_container_width=True
                    )

                zip_buffer = io.BytesIO()

                with zipfile.ZipFile(
                    zip_buffer,
                    "w",
                    zipfile.ZIP_DEFLATED
                ) as zip_file:

                    for file_path in created_files:
                        zip_file.write(
                            file_path,
                            arcname=os.path.basename(file_path)
                        )

                zip_buffer.seek(0)

                st.divider()

                st.download_button(
                    label="📦 Download All MRF Files (ZIP)",
                    data=zip_buffer,
                    file_name="MRF_All_Files.zip",
                    mime="application/zip",
                    use_container_width=True
                )

        except Exception as e:
            st.error(
                "❌ An error occurred while generating the MRF files."
            )
            st.exception(e)


st.markdown(
    '<div class="footer-note">MRF Automation • Excel processing workspace</div>',
    unsafe_allow_html=True
)